import os
import sys
from io import BytesIO

from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

from import_report_utils import import_report_workbook
from import_routes import _teacher_import_error


report = {
    "kind": "enseignants",
    "created": 1,
    "skipped": 2,
    "errors": [
        _teacher_import_error(4, "Département", "Le département est absent ou introuvable.", "Choisissez un département existant ou indiquez son code exact."),
        _teacher_import_error(7, "Heures dues", "La valeur n’est pas un nombre entier.", "Saisissez uniquement un nombre entier, par exemple 18."),
    ],
    "has_more": False,
    "generated_at": "06/09/2026 09:00",
}

workbook = import_report_workbook(report)
loaded = load_workbook(BytesIO(workbook.getvalue()), data_only=True)
rows = list(loaded.active.iter_rows(values_only=True))
assert rows[6][:4] == ("Ligne", "Champ", "Problème", "Correction recommandée")
assert rows[7][:4] == (4, "Département", "Le département est absent ou introuvable.", "Choisissez un département existant ou indiquez son code exact.")
assert rows[8][0] == 7 and rows[8][1] == "Heures dues"
loaded.close()

html = open(os.path.join(BASE_DIR, "templates/import_report.html"), encoding="utf-8").read()
pdf = open(os.path.join(BASE_DIR, "templates/pdf/import_report_pdf.html"), encoding="utf-8").read()
assert "Correction recommandée" in html
assert "error.cause" in html and "error.correction" in html
assert "error.field" in pdf and "error.correction" in pdf

print("TEACHER_IMPORT_REPORT_FEATURE_TEST_OK")
